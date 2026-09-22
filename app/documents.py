from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


class DocumentError(ValueError):
    """A user-correctable document processing error."""


SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md", ".py", ".js", ".ts", ".java", ".json", ".csv"}
DOCUMENT_EXTENSIONS = {".docx", ".pdf"}


@dataclass
class Block:
    block_id: str
    text: str
    kind: str = "paragraph"
    page: int | None = None
    target: Any = field(default=None, repr=False)


@dataclass
class ParsedDocument:
    path: Path
    extension: str
    blocks: list[Block]
    title: str
    source_document: Any = field(default=None, repr=False)

    @property
    def text(self) -> str:
        return "\n\n".join(block.text for block in self.blocks if block.text.strip())


def extension_for(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError("This file type is not supported.")
    return extension


def _iter_block_items(parent: DocumentType | _Cell):
    if isinstance(parent, DocumentType):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        raise TypeError(f"unsupported document parent: {type(parent)!r}")

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _parse_docx(path: Path) -> ParsedDocument:
    document = Document(path)
    blocks: list[Block] = []
    paragraph_number = 0
    table_number = 0

    for item in _iter_block_items(document):
        if isinstance(item, Paragraph):
            paragraph_number += 1
            blocks.append(
                Block(
                    block_id=f"p{paragraph_number:05d}",
                    text=item.text,
                    kind="paragraph",
                    target=item,
                )
            )
        elif isinstance(item, Table):
            table_number += 1
            for row_number, row in enumerate(item.rows):
                for cell_number, cell in enumerate(row.cells):
                    for paragraph_number_in_cell, paragraph in enumerate(cell.paragraphs):
                        blocks.append(
                            Block(
                                block_id=(
                                    f"t{table_number:04d}"
                                    f"r{row_number:04d}c{cell_number:04d}"
                                    f"p{paragraph_number_in_cell:04d}"
                                ),
                                text=paragraph.text,
                                kind="table_cell",
                                target=paragraph,
                            )
                        )

    if sum(len(block.text) for block in blocks) == 0:
        raise DocumentError("The DOCX does not contain readable text.")
    return ParsedDocument(path, ".docx", blocks, path.name, document)


def _parse_pdf(path: Path, max_pages: int) -> ParsedDocument:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pypdf exposes different exception types by version.
        raise DocumentError("The PDF could not be opened.") from exc
    if len(reader.pages) > max_pages:
        raise DocumentError(f"PDFs are limited to {max_pages} pages.")

    blocks: list[Block] = []
    total_text = 0
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise DocumentError(f"Could not extract text from PDF page {page_number}.") from exc
        text = text.strip()
        total_text += len(text)
        if text:
            blocks.append(
                Block(
                    block_id=f"page{page_number:04d}",
                    text=text,
                    kind="pdf_page",
                    page=page_number,
                )
            )

    if total_text < 40:
        raise DocumentError("This PDF appears to be scanned or image-only; OCR is not enabled.")
    return ParsedDocument(path, ".pdf", blocks, path.name)


def _parse_text(path: Path, extension: str) -> ParsedDocument:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise DocumentError("The text file could not be read.") from exc
    if not text.strip():
        raise DocumentError("The file does not contain readable text.")
    blocks = [Block(f"line{i:06d}", line) for i, line in enumerate(text.splitlines(), start=1)]
    return ParsedDocument(path, extension, blocks, path.name)


def parse_document(path: Path, *, max_pdf_pages: int = 50, max_docx_chars: int = 200_000) -> ParsedDocument:
    extension = path.suffix.lower()
    if extension == ".docx":
        parsed = _parse_docx(path)
    elif extension == ".pdf":
        parsed = _parse_pdf(path, max_pdf_pages)
    elif extension in SUPPORTED_EXTENSIONS:
        parsed = _parse_text(path, extension)
    else:
        raise DocumentError("This file type is not supported.")

    if extension == ".docx" and len(parsed.text) > max_docx_chars:
        raise DocumentError(f"DOCX files are limited to about {max_docx_chars:,} extracted characters.")
    return parsed


def make_chunks(blocks: list[Block], max_chars: int = 12_000) -> list[list[Block]]:
    chunks: list[list[Block]] = []
    current: list[Block] = []
    current_size = 0
    for block in blocks:
        block_size = len(block.text) + len(block.block_id) + 20
        if current and current_size + block_size > max_chars:
            chunks.append(current)
            current = []
            current_size = 0
        current.append(block)
        current_size += block_size
    if current:
        chunks.append(current)
    return chunks


def format_chunk(chunk: list[Block]) -> str:
    return "\n\n".join(f"[BLOCK {block.block_id}]\n{block.text}" for block in chunk)


def source_label(block: Block) -> str:
    if block.page is not None:
        return f"Page {block.page}"
    if block.kind == "table_cell":
        return f"Table cell {block.block_id}"
    if block.block_id.startswith("line"):
        return f"Line {block.block_id.removeprefix('line')}"
    return f"Paragraph {block.block_id.removeprefix('p')}"


def format_context_blocks(blocks: list[Block]) -> str:
    return "\n\n".join(f"[SOURCE: {source_label(block)}]\n{block.text}" for block in blocks if block.text.strip())


def select_context(parsed: ParsedDocument, question: str, max_chars: int = 30_000) -> str:
    full_text = parsed.text
    if len(full_text) <= max_chars:
        return format_context_blocks(parsed.blocks)

    words = {
        word.lower()
        for word in re.findall(r"[\w\u0900-\u097F]{3,}", question)
        if word.lower() not in {"what", "where", "when", "which", "that", "this", "the", "and"}
    }
    chunks = make_chunks(parsed.blocks, max_chars=max(1, max_chars // 2))
    scored: list[tuple[int, int, str]] = []
    for index, chunk in enumerate(chunks):
        text = format_context_blocks(chunk)
        score = sum(text.lower().count(word) for word in words)
        scored.append((score, index, text))
    selected = sorted(scored, key=lambda item: (-item[0], item[1]))[: max(1, max_chars // 5_000)]
    selected.sort(key=lambda item: item[1])
    return "\n\n".join(item[2] for item in selected)[:max_chars]


def _replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def build_corrected_docx(parsed: ParsedDocument, corrections: dict[str, str], output_path: Path) -> None:
    if parsed.extension == ".docx":
        for block in parsed.blocks:
            if block.block_id in corrections and block.target is not None:
                _replace_paragraph_text(block.target, corrections[block.block_id])
        parsed.source_document.save(output_path)
        return

    output = Document()
    for index, block in enumerate(parsed.blocks):
        output.add_paragraph(corrections.get(block.block_id, block.text))
        if block.page is not None and index < len(parsed.blocks) - 1:
            next_page = parsed.blocks[index + 1].page
            if next_page is not None and next_page != block.page:
                output.add_page_break()
    output.save(output_path)
