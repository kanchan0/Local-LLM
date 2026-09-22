from pathlib import Path

import pytest
from docx import Document

from app.correction import correct_document
from app.documents import build_corrected_docx, make_chunks, parse_document, select_context


def test_docx_round_trip_preserves_basic_structure(tmp_path: Path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("A heading", level=1)
    document.add_paragraph("This sentence have an error.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First cell"
    table.cell(0, 1).text = "Second cell"
    document.save(source)

    parsed = parse_document(source)
    assert any(block.text == "A heading" for block in parsed.blocks)
    assert any(block.kind == "table_cell" for block in parsed.blocks)

    output = tmp_path / "corrected.docx"
    corrections = {block.block_id: block.text.replace("have", "has") for block in parsed.blocks}
    build_corrected_docx(parsed, corrections, output)

    result = parse_document(output)
    assert "This sentence has an error." in result.text
    assert "First cell" in result.text


def test_chunking_keeps_block_order():
    from app.documents import Block

    blocks = [Block(f"p{i}", "x" * 20) for i in range(10)]
    chunks = make_chunks(blocks, max_chars=70)
    flattened = [block.block_id for chunk in chunks for block in chunk]
    assert flattened == [block.block_id for block in blocks]


def test_large_context_selects_relevant_chunks():
    from app.documents import Block, ParsedDocument

    parsed = ParsedDocument(
        Path("example.txt"),
        ".txt",
        [Block("a", "unrelated " * 100), Block("b", "database migration details " * 100)],
        "example.txt",
    )
    context = select_context(parsed, "What are the database migration details?", max_chars=500)
    assert "database migration" in context
    assert "[SOURCE: Paragraph b]" in context


@pytest.mark.asyncio
async def test_correction_service_generates_docx(tmp_path: Path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("This sentence have an error.")
    document.save(source)

    class FakeOllama:
        async def correct_chunk(self, chunk):
            return {
                "blocks": [
                    {
                        "id": block.block_id,
                        "corrected_text": block.text.replace("have", "has"),
                        "notes": ["Fixed subject-verb agreement."],
                    }
                    for block in chunk
                ]
            }

    output = tmp_path / "corrected.docx"
    result = await correct_document(
        source,
        output,
        "source.docx",
        FakeOllama(),
        max_pdf_pages=50,
        max_docx_chars=200_000,
    )
    assert result.output_name == "source_corrected.docx"
    assert output.exists()
    assert "This sentence has an error." in parse_document(output).text
